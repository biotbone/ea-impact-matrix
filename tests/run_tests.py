#!/usr/bin/env python3
"""Test runner for the ea-impact-matrix skill.

Layers
  A  Acceptance: the config produced from the scenario documents handles every planted trap.
  B  Workbook: the recalculated workbook agrees with an independent recomputation (the oracle).
  C  Tool: behaviour of the script (flags, exit codes, validation, compatibility, public-copy hygiene).

Usage
  python3 tests/run_tests.py            run everything, print a table, exit 1 on any failure
  python3 tests/run_tests.py --report   also write tests/TEST_REPORT.md

Needs python3, openpyxl and LibreOffice (soffice) on the path. Nothing is written outside a temp folder
except TEST_REPORT.md when --report is given.
"""
import argparse, copy, datetime, importlib.util, json, os, re, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SCRIPT = os.path.join(ROOT, "scripts", "build_impact_matrix.py")
SCEN = os.path.join(HERE, "scenario", "output", "portfolio.json")
LEGACY = os.path.join(HERE, "fixtures", "legacy_config.json")
EXAMPLE = os.path.join(ROOT, "examples", "portfolio.example.json")
TMP = tempfile.mkdtemp(prefix="eaim_")

DIMS = ["business_value", "cost_optimization", "time_criticality", "risk_reduction",
        "security_compliance", "delivery_confidence", "strategic_alignment"]
RUN_COST = [0.10, 0.35, 0.15, 0.10, 0.05, 0.15, 0.10]
ROWS = []  # (id, reqs, layer, name, ok, detail)
TESTS = []


def test(tid, reqs, layer, name):
    def deco(fn):
        TESTS.append((tid, reqs, layer, name, fn))
        return fn
    return deco


def load(path):
    with open(path) as f:
        return json.load(f)


def save(cfg, name):
    p = os.path.join(TMP, name)
    with open(p, "w") as f:
        json.dump(cfg, f)
    return p


def run(*args):
    r = subprocess.run([sys.executable, SCRIPT, *args], capture_output=True, text=True)
    return r.returncode, r.stdout, r.stderr


def build_recalc(cfg_path, tag):
    out = os.path.join(TMP, tag + ".xlsx")
    rc, so, se = run(cfg_path, "-o", out)
    if rc != 0:
        raise AssertionError(f"build failed rc={rc}: {se.strip()[:200]}")
    rd = os.path.join(TMP, tag + "_rc")
    os.makedirs(rd, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", rd, out],
                   capture_output=True, text=True, timeout=180)
    return os.path.join(rd, tag + ".xlsx")


def oracle(cfg):
    """Independent recomputation: weighted score, rank, score per $100k."""
    w = [cfg["weights"][k] for k in DIMS]
    rows = []
    for p in cfg["projects"]:
        if not p.get("scores"):
            continue
        s = sum(a * b for a, b in zip(w, (p["scores"][k] for k in DIMS)))
        rows.append(dict(id=p["id"], score=s, cost=p["cost"], per100=s / (p["cost"] / 1e5) if p["cost"] else None))
    order = sorted(rows, key=lambda r: -r["score"])
    for i, r in enumerate(order, 1):
        r["rank"] = i
    return {r["id"]: r for r in rows}


_cache = {}


def scenario_wb():
    if "wb" not in _cache:
        import openpyxl
        path = build_recalc(SCEN, "scenario")
        _cache["wb"] = openpyxl.load_workbook(path, data_only=True)
        _cache["path"] = path
    return _cache["wb"]


def assert_(cond, msg):
    if not cond:
        raise AssertionError(msg)


def proj(cfg, pid):
    return next(p for p in cfg["projects"] if p["id"] == pid)


def text_of(obj):
    return json.dumps(obj, ensure_ascii=False)


# ------------------------------------------------------------------ A: acceptance on the scenario config

@test("A01", "R16", "A", "Scenario config passes --check with no errors and exactly the 2 expected warnings")
def _():
    rc, so, se = run(SCEN, "--check")
    assert_(rc == 0, f"rc={rc}\n{so}")
    warns = [l for l in so.splitlines() if l.startswith("WARN")]
    assert_(len(warns) == 2 and "Ingest volume" in so and "run cost per month" in so, so)
    return "0 errors, 2 warnings (the two KPI baselines logged as open questions)"


@test("A02", "R01", "A", "Lens is run_cost, weights equal the preset, approver and date recorded from LENS.md")
def _():
    c = load(SCEN); pf = c["portfolio"]
    assert_(pf["lens"] == "run_cost", pf["lens"])
    assert_([c["weights"][k] for k in DIMS] == RUN_COST, c["weights"])
    assert_("Whitcombe" in pf["lens_approved_by"] and pf["lens_approved_date"] == "2026-03-02", pf)
    assert_(pf["next_review"] == "2026-06-01" and pf["risk_threshold"] == 9, pf)


@test("A03", "R15", "A", "Five initiatives; P5 (idea with no estimate) is not scored")
def _():
    c = load(SCEN)
    assert_([p["id"] for p in c["projects"]] == ["P1", "P2", "P3", "P4", "P5"], "ids")
    assert_(not proj(c, "P5").get("scores"), "P5 has scores")
    assert_(proj(c, "P5")["wave"] == "Hold", "P5 not held")


@test("A04", "R06", "A", "P4 cost uses the line items (220,000) and the 205,000 deck total is called out")
def _():
    p = proj(load(SCEN), "P4")
    assert_(p["cost"] == 220000, p["cost"])
    assert_("205,000" in p["cost_note"] and "220,000" in p["cost_note"], p["cost_note"])


@test("A05", "R11", "A", "P4 security conflict with SEC-07 lowers the score and fails the standards gate")
def _():
    p = proj(load(SCEN), "P4")
    assert_(p["scores"]["security_compliance"] <= 2, p["scores"])
    assert_("SEC-07" in p["assessments"]["security_compliance"]["assessment"], "no SEC-07")
    assert_(p["gates"][1][0] == "Not met", p["gates"][1])
    assert_(p["scores"]["strategic_alignment"] <= 2, "alignment not penalised")


@test("A06", "R05", "A", "Ranking is P2, P1, P3, P4 under the run-cost lens")
def _():
    o = oracle(load(SCEN))
    order = [k for k, _ in sorted(o.items(), key=lambda kv: kv[1]["rank"])]
    assert_(order == ["P2", "P1", "P3", "P4"], order)


@test("A07", "R09", "A", "Wave 1 is P1 and P2, costs 144,000 and fits the 150,000 first tranche; P4 and P5 held")
def _():
    c = load(SCEN)
    w = {r[0]: r for r in c["waves"]}
    assert_(set(x.strip() for x in w["Wave 1"][1].split(",")) == {"P1", "P2"}, w["Wave 1"])
    assert_(w["Wave 1"][3] == proj(c, "P1")["cost"] + proj(c, "P2")["cost"] == 144000, w["Wave 1"][3])
    assert_(w["Wave 1"][3] <= 150000, "over tranche")
    assert_(set(x.strip() for x in w["Hold"][1].split(",")) == {"P4", "P5"}, w["Hold"])
    assert_(w["Wave 2"][1].strip() == "P3", w["Wave 2"])


@test("A08", "R10", "A", "Cross-dependencies include the shared-architect conflict, P1 enabling P3 and P4, and P3 with P4")
def _():
    deps = [tuple(r[:3]) for r in load(SCEN)["cross_dependencies"]]
    for want in [("P2", "P4", "Capacity"), ("P1", "P3", "Enables"), ("P1", "P4", "Enables"), ("P3", "P4", "Commercial")]:
        assert_(want in deps, f"missing {want}")


@test("A09", "R07", "A", "No unevidenced saving is claimed: P4 zero, P3 below the vendor 45%, P2 matches the baseline")
def _():
    s = load(SCEN)["savings"]
    assert_(s["P4"]["expected_saving"] == 0, s["P4"])
    assert_(s["P3"]["expected_saving"] <= 4500 * 12, s["P3"])
    assert_(s["P3"]["expected_saving"] < 0.45 * 120000, "uses the vendor claim")
    assert_(s["P2"]["expected_saving"] == (9400 - 3900) * 12, s["P2"])
    assert_("P1" not in s or s["P1"]["expected_saving"] == 0, "enabler claims saving")


@test("A10", "R08", "A", "All scored initiatives have distinct score profiles")
def _():
    profs = [tuple(p["scores"][k] for k in DIMS) for p in load(SCEN)["projects"] if p.get("scores")]
    assert_(len(set(profs)) == len(profs), "duplicate profile")


@test("A11", "R14", "A", "Open questions capture every planted gap")
def _():
    t = text_of(load(SCEN)["open_questions"])
    for token in ["cutover window", "retained", "tags", "ingest", "2026-09-16", "baseline", "205,000", "SEC-07",
                  "1,100", "business case", "second tranche"]:
        assert_(token.lower() in t.lower(), f"no open question mentions '{token}'")


@test("A12", "R02", "A", "Document register shows the real gaps: no P5 estimate, no P4 cost baseline")
def _():
    d = {(r[0], r[1]): r[5] for r in load(SCEN)["documents"]}
    assert_(d[("ESTIMATE", "P5")] == "Not available", d.get(("ESTIMATE", "P5")))
    assert_(d[("COST_BASELINE", "P4")] == "Requested", d.get(("COST_BASELINE", "P4")))
    for code in ["TARGET_STATE", "LENS", "FUNDING"]:
        assert_(d[(code, "ALL")] == "Received", code)


@test("A13", "R03", "A", "Every score on a scored initiative has an assessment, an owner and evidence")
def _():
    for p in load(SCEN)["projects"]:
        if not p.get("scores"):
            continue
        for k in DIMS:
            a = p["assessments"][k]
            assert_(a["assessment"].strip() and a["evidence"].strip() and a["owner"].strip(), f"{p['id']} {k}")


@test("A14", "R11", "A", "P4 capacity trap: assessment names the 1,100 hours against about 800 available")
def _():
    a = proj(load(SCEN), "P4")["assessments"]["delivery_confidence"]["assessment"]
    assert_("1,100" in a and "800" in a, a)
    assert_(proj(load(SCEN), "P4")["scores"]["delivery_confidence"] <= 2, "confidence too high")


@test("A15", "R12", "A", "Every risk at or above the threshold has a named owner and a response")
def _():
    L = {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}
    for p in load(SCEN)["projects"]:
        for r in p.get("risks", []):
            if L[r[1]] * L[r[2]] >= 9:
                assert_(r[4].strip() and len(r) > 5 and r[5], f"{p['id']} {r[0]}")


@test("A16", "R13", "A", "Decision log holds a recommendation for every initiative, with a named decider")
def _():
    d = {r[0]: r for r in load(SCEN)["decisions"]}
    assert_(set(d) == {"P1", "P2", "P3", "P4", "P5"}, set(d))
    assert_(d["P4"][1] == "Held" and d["P3"][1] == "Deferred" and d["P2"][1] == "Approved with conditions", d)


# ------------------------------------------------------------------ B: workbook against the oracle

@test("B01", "R19", "B", "Recalculated scenario workbook has no formula errors")
def _():
    wb = scenario_wb()
    bad = [(s.title, c.coordinate, c.value) for s in wb for row in s.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("#")]
    assert_(not bad, bad[:5])


@test("B02", "R04", "B", "Weighted score D27 on each project sheet equals the oracle")
def _():
    wb = scenario_wb(); o = oracle(load(SCEN))
    for pid, r in o.items():
        v = wb[pid]["D27"].value
        assert_(abs(v - r["score"]) < 1e-9, f"{pid} {v} vs {r['score']}")


@test("B03", "R05", "B", "Summary ranks and score per $100k equal the oracle")
def _():
    wb = scenario_wb(); o = oracle(load(SCEN)); ws = wb["02 Portfolio Summary"]
    seen = 0
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        if row[0] in o:
            r = o[row[0]]; seen += 1
            assert_(row[5] == r["rank"], f"{row[0]} rank {row[5]} vs {r['rank']}")
            assert_(abs(row[6] - r["per100"]) < 1e-6, f"{row[0]} per100 {row[6]} vs {r['per100']}")
    assert_(seen == 4, f"found {seen} scored rows")


@test("B04", "R04", "B", "Fixed cell addresses carry the config values (B5, B6, B10, B11, B14, B20:B26)")
def _():
    wb = scenario_wb(); c = load(SCEN)
    for p in c["projects"]:
        ws = wb[p["sheet"]]
        assert_(ws["B5"].value == p["id"] and ws["B6"].value == p["name"], f"{p['id']} id or name")
        if p.get("scores"):
            assert_(ws["B10"].value == p["cost"] and ws["B11"].value == p["months"] and ws["B14"].value == p["wave"], p["id"])
            for i, k in enumerate(DIMS):
                assert_(ws.cell(row=20 + i, column=2).value == p["scores"][k], f"{p['id']} {k}")


@test("B05", "R20", "B", "Workbook has the expected sheets in order")
def _():
    n = scenario_wb().sheetnames
    assert_(n == ["00 Read Me", "01 Scoring Model", "02 Portfolio Summary", "03 Sequencing", "04 Risk Register",
                  "05 Open Questions", "06 Savings Case", "07 Decision Log", "P1", "P2", "P3", "P4", "P5"], n)


@test("B06", "R07", "B", "Savings sheet: P2 payback is cost over annual saving; initiatives with no baseline show no payback")
def _():
    ws = scenario_wb()["06 Savings Case"]; got = {}
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        if row[0] in ("P1", "P2", "P3", "P4"):
            got[row[0]] = row[6]
    assert_(abs(got["P2"] - 96000 / 66000 * 12) < 1e-6, got)
    assert_(got["P1"] == "Awaiting baseline" and got["P4"] == "Awaiting baseline", got)


@test("B07", "R12", "B", "Risk register lists every configured risk")
def _():
    ws = scenario_wb()["04 Risk Register"]
    n_cfg = sum(len(p.get("risks", [])) for p in load(SCEN)["projects"])
    n_ws = sum(1 for row in ws.iter_rows(values_only=True) if row[0] in {"P1", "P2", "P3", "P4"} or
               (row[1] in {"P1", "P2", "P3", "P4"}))
    assert_(n_ws >= n_cfg, f"{n_ws} rows for {n_cfg} risks")


@test("B08", "R15", "B", "An unscored initiative (P5) has no score or rank and is not shown as zero cost")
def _():
    wb = scenario_wb(); ws = wb["02 Portfolio Summary"]
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        if row[0] == "P5":
            assert_(row[4] in (None, "") and row[5] in (None, ""), f"P5 score/rank {row[4]} {row[5]}")
            assert_(row[2] in (None, "", "Not scored", "TBC"), f"P5 cost shown as {row[2]!r}")
            return
    raise AssertionError("P5 row not found")


@test("B09", "R05", "B", "Data validation lists are present (risk response, gate status, decision)")
def _():
    import openpyxl
    wb = openpyxl.load_workbook(_cache["path"])
    n = sum(len(s.data_validations.dataValidation) for s in wb)
    assert_(n >= 3, f"{n} validation rules")


# ------------------------------------------------------------------ C: tool behaviour

@test("C01", "R20", "C", "Example config: 0 errors, 1 warning; recalculated D27 is 3.5")
def _():
    rc, so, se = run(EXAMPLE, "--check")
    assert_(rc == 0 and "0 error(s), 1 warning(s)" in so, so)
    import openpyxl
    wb = openpyxl.load_workbook(build_recalc(EXAMPLE, "example"), data_only=True)
    assert_(abs(wb["P1"]["D27"].value - 3.5) < 1e-9, wb["P1"]["D27"].value)


@test("C02", "R20", "C", "Blank template builds and --dump-config output passes --check without errors")
def _():
    out = os.path.join(TMP, "tpl.xlsx")
    rc, so, se = run("--template", "-o", out)
    assert_(rc == 0 and os.path.exists(out), se)
    d = os.path.join(TMP, "dump.json")
    rc, so, se = run("--dump-config", d)
    assert_(rc == 0 and os.path.exists(d), se)
    rc, so, se = run(d, "--check")
    assert_(rc in (0, 1) and "0 error(s)" in so, so)


@test("C03", "R02", "C", "Catalog has 13 document types, 5 required, and --intake lists every one")
def _():
    spec = importlib.util.spec_from_file_location("eaim", SCRIPT)
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    cat = m.DOC_CATALOG
    assert_(len(cat) == 13, len(cat))
    assert_(sum(1 for c in cat if c[1] == "Required") == 5, "required count")
    rc, so, se = run("--intake")
    assert_(rc == 0, se)
    for c in cat:
        assert_(c[0] in so, f"{c[0]} not in --intake")


@test("C04", "R02", "C", "references/required-documents.md matches what --intake prints")
def _():
    rc, so, se = run("--intake")
    with open(os.path.join(ROOT, "references", "required-documents.md")) as f:
        ref = f.read()
    assert_(so.strip() == ref.strip(), "required-documents.md is out of date: regenerate it from --intake")


def mutate(fn, name):
    c = load(SCEN); fn(c); return save(c, name)


@test("C05", "R17", "C", "Score outside 1 to 5 is an error, exit 2, and no workbook is written")
def _():
    p = mutate(lambda c: proj(c, "P1")["scores"].update(business_value=6), "badscore.json")
    out = os.path.join(TMP, "badscore.xlsx")
    rc, so, se = run(p, "-o", out)
    assert_(rc == 2 and not os.path.exists(out), f"rc={rc} exists={os.path.exists(out)}")


@test("C06", "R16", "C", "--strict turns warnings into exit 1 and blocks the build; clean config passes strict")
def _():
    out = os.path.join(TMP, "strict.xlsx")
    rc, so, se = run(SCEN, "--strict", "-o", out)
    assert_(rc == 1 and not os.path.exists(out), f"rc={rc}")
    def fix(c):
        for p in c["projects"]:
            for k in p.get("kpis", []):
                if k[1] in ("TBC, not in finance export", "TBC"):
                    k[1] = "Measured, 10,000"
    p = mutate(fix, "clean.json")
    rc, so, se = run(p, "--check", "--strict")
    assert_(rc == 0, f"rc={rc}\n{so}")


@test("C07", "R17", "C", "Weights that do not sum to 1 are an error")
def _():
    p = mutate(lambda c: c["weights"].update(business_value=0.2), "w.json")
    rc, so, se = run(p, "--check")
    assert_(rc == 2 and "weights sum" in so, so)


@test("C08", "R17", "C", "Bad gate status, unknown document code, wrong document project and bad decision are errors")
def _():
    cases = {
        "gate status": lambda c: proj(c, "P1")["gates"].__setitem__(0, ["Maybe", "x", "y"]),
        "unknown document code": lambda c: c["documents"].append(["NOPE", "ALL", "x", "", "x", "Received"]),
        "is not ALL or a project id": lambda c: c["documents"].append(["PROPOSAL", "P9", "x", "", "x", "Received"]),
        "is not one of Approved": lambda c: c["decisions"].append(["P1", "Maybe", "", "x", "", ""]),
    }
    for needle, fn in cases.items():
        rc, so, se = run(mutate(fn, "case.json"), "--check")
        assert_(rc == 2 and needle in so, f"{needle}: rc={rc}\n{so}")


@test("C09", "R08", "C", "Identical score profiles are flagged")
def _():
    p = mutate(lambda c: proj(c, "P3").__setitem__("scores", dict(proj(c, "P1")["scores"])), "dup.json")
    rc, so, se = run(p, "--check")
    assert_("identical score profile" in so and "P1, P3" in so, so)


@test("C10", "R12", "C", "risk_threshold changes what counts as a red risk")
def _():
    def add(c, thr):
        c["portfolio"]["risk_threshold"] = thr
        proj(c, "P1")["risks"].append(["Ownerless high risk", "High", "High", "Fix it.", ""])
    p9 = mutate(lambda c: add(c, 9), "t9.json")
    p12 = mutate(lambda c: add(c, 12), "t12.json")
    rc, so, se = run(p9, "--check")
    assert_("red risk 'Ownerless high risk' has no owner" in so, so)
    rc, so, se = run(p12, "--check")
    assert_("Ownerless high risk" not in so, so)
    rc, so, se = run(mutate(lambda c: add(c, 40), "t40.json"), "--check")
    assert_(rc == 2 and "risk_threshold" in so, so)


@test("C11", "R12", "C", "Residual above initial score warns; invalid residual and invalid response are errors")
def _():
    def worse(c):
        proj(c, "P1")["risks"][1] = ["Low risk", "Low", "Low", "x", "Owner", "Mitigate", "Critical", "Critical"]
    rc, so, se = run(mutate(worse, "res.json"), "--check")
    assert_("residual is higher" in so and rc == 0, so)
    def badresp(c):
        proj(c, "P1")["risks"][1] = ["R", "Low", "Low", "x", "Owner", "Ignore"]
    rc, so, se = run(mutate(badresp, "resp.json"), "--check")
    assert_(rc == 2 and "response must be one of" in so, so)
    def badres(c):
        proj(c, "P1")["risks"][1] = ["R", "Low", "Low", "x", "Owner", "Mitigate", "Huge", "Low"]
    rc, so, se = run(mutate(badres, "res2.json"), "--check")
    assert_(rc == 2 and "invalid residual" in so, so)


@test("C12", "R18", "C", "Legacy config (no new optional keys) builds and gives the same scores as the oracle")
def _():
    c = load(LEGACY)
    for k in ("documents", "decisions"):
        assert_(k not in c, f"fixture should not have {k}")
    rc, so, se = run(LEGACY, "--check")
    assert_(rc == 0, so)
    import openpyxl
    wb = openpyxl.load_workbook(build_recalc(LEGACY, "legacy"), data_only=True)
    for pid, r in oracle(c).items():
        assert_(abs(wb[pid]["D27"].value - r["score"]) < 1e-9, pid)


@test("C13", "R19", "C", "Two builds from the same config give identical cell contents")
def _():
    import openpyxl
    def dump(path):
        wb = openpyxl.load_workbook(path)
        return {(s.title, c.coordinate): c.value for s in wb for row in s.iter_rows() for c in row if c.value is not None}
    a, b = os.path.join(TMP, "d1.xlsx"), os.path.join(TMP, "d2.xlsx")
    run(SCEN, "-o", a); run(SCEN, "-o", b)
    assert_(dump(a) == dump(b), "cell contents differ between builds")


@test("C14", "R16", "C", "Every flag documented in SKILL.md exists in the script")
def _():
    with open(os.path.join(ROOT, "SKILL.md")) as f:
        t = f.read()
    flags = set()
    for line in t.splitlines():
        if "build_impact_matrix" in line:
            flags |= set(re.findall(r"(?<![\w-])(--[a-z][a-z-]+)", line))
    rc, so, se = run("--help")
    missing = [f for f in flags if f not in so]
    assert_(not missing, missing)


@test("C15", "R21", "C", "Public copy contains no client or employer names")
def _():
    banned = ["Advantage Solutions", "Reply S", "Reply Group", "Thiagovault", "Obsidian"]
    hits = []
    for dp, dn, fn in os.walk(ROOT):
        dn[:] = [d for d in dn if d not in ("__pycache__", ".git")]
        for f in fn:
            if f.endswith((".xlsx", ".pyc")) or f in ("TEST_REPORT.md",) or f == "run_tests.py":
                continue
            with open(os.path.join(dp, f), errors="ignore") as fh:
                t = fh.read()
            hits += [(os.path.relpath(os.path.join(dp, f), ROOT), b) for b in banned if b in t]
    assert_(not hits, hits)
    import openpyxl
    for x in [os.path.join(ROOT, "EA_Impact_Matrix_TEMPLATE.xlsx"), _cache["path"]]:
        wb = openpyxl.load_workbook(x)
        for s in wb:
            for row in s.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        assert_(not any(b in c.value for b in banned), f"{x} {s.title} {c.coordinate}")


@test("C16", "R21", "C", "No em dash or en dash in any text file or workbook of the public copy")
def _():
    import openpyxl
    bad = []
    for dp, dn, fn in os.walk(ROOT):
        dn[:] = [d for d in dn if d not in ("__pycache__", ".git")]
        for f in fn:
            p = os.path.join(dp, f)
            if f.endswith(".pyc"):
                continue
            if f.endswith(".xlsx"):
                wb = openpyxl.load_workbook(p)
                for s in wb:
                    for row in s.iter_rows():
                        for c in row:
                            if isinstance(c.value, str) and (chr(0x2014) in c.value or chr(0x2013) in c.value):
                                bad.append((f, s.title, c.coordinate))
            else:
                with open(p, errors="ignore") as fh:
                    t = fh.read()
                if chr(0x2014) in t or chr(0x2013) in t:
                    bad.append(os.path.relpath(p, ROOT))
    assert_(not bad, bad[:8])


@test("C17", "R21", "C", ".gitignore excludes bytecode and a licence file is present")
def _():
    with open(os.path.join(ROOT, ".gitignore")) as f:
        g = f.read()
    assert_("__pycache__/" in g and "*.pyc" in g, g)
    assert_(os.path.exists(os.path.join(ROOT, "LICENSE")), "no LICENSE")


# ------------------------------------------------------------------ driver

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args()
    for tid, reqs, layer, name, fn in TESTS:
        try:
            d = fn()
            ROWS.append((tid, reqs, layer, name, True, d or ""))
        except AssertionError as e:
            ROWS.append((tid, reqs, layer, name, False, str(e).replace("\n", " | ")[:300]))
        except Exception as e:  # runner or environment problem, still a failure
            ROWS.append((tid, reqs, layer, name, False, f"{type(e).__name__}: {e}"[:300]))
        print(("PASS " if ROWS[-1][4] else "FAIL ") + tid + "  " + name)
        if not ROWS[-1][4]:
            print("      " + ROWS[-1][5])
    ok = sum(1 for r in ROWS if r[4])
    print(f"\n{ok}/{len(ROWS)} passed")
    if a.report:
        write_report(ok)
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(0 if ok == len(ROWS) else 1)


def write_report(ok):
    L = ["# Test report: ea-impact-matrix", "",
         f"Run date: {datetime.date.today().isoformat()}. Result: **{ok} of {len(ROWS)} tests passed.**", "",
         "Generated by `tests/run_tests.py --report`. Test design and traceability are in `TEST_PLAN.md`.", "",
         "| ID | Layer | Requirements | Test | Result | Detail |", "|---|---|---|---|---|---|"]
    for tid, reqs, layer, name, passed, d in ROWS:
        L.append(f"| {tid} | {layer} | {reqs} | {name} | {'Pass' if passed else 'FAIL'} | {d.replace('|', '/')} |")
    L += ["", "## Reading the result", "",
          "- Layer A checks that the configuration produced from the scenario documents handles every planted trap.",
          "- Layer B checks the recalculated workbook against a separate recomputation of the scores, ranks and payback.",
          "- Layer C checks the script itself: flags, exit codes, validation, backward compatibility and public-copy hygiene.", ""]
    L += ["## Defects found during the first run", "",
          "| # | Where | Finding | Resolution |", "|---|---|---|---|",
          "| 1 | Tool | An initiative with no scores (held, for example) showed cost 0, duration 0 and ratings 0 on 02 Portfolio Summary, which reads as a free project | Fixed. Blank cost and duration stay empty on the project sheet and the summary shows empty cells (B08) |",
          "| 2 | Test A11 | Looked for the word retention, the open question says retained | Test corrected |",
          "| 3 | Test C14 | Treated LibreOffice flags in SKILL.md as script flags | Test now checks only lines that call the script |",
          "| 4 | Test C16 | The runner itself contained dash characters inside its own check | Rewritten with character codes |", "",
          "## Observations, not defects", "",
          "- The portfolio payback on 06 Savings Case divides the cost of every initiative, including enablers and held ones, by the savings that exist. It is correct arithmetic but can read pessimistically. The per-project rows are the ones to quote.",
          "- The scenario author, config author and runner are the same agent. A human architect should review the scenario and the scores.",
          "- Recalculation uses LibreOffice. Open the shipped workbooks once in Excel before relying on them.", ""]
    with open(os.path.join(HERE, "TEST_REPORT.md"), "w") as f:
        f.write("\n".join(L))


if __name__ == "__main__":
    main()
